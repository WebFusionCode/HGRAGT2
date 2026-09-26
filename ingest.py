"""Load a manifest-driven, mixed-format synthetic knowledge corpus."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any


DEFAULT_CORPUS_DIR = Path(__file__).parent / "data" / "synthetic"
MAX_CHUNK_CHARS = 1_100
CHUNK_OVERLAP_CHARS = 140

_PII_PATTERNS = (
    re.compile(r"\b(?:MRN|medical record(?: number| #)?|patient ID)\s*[:#=-]?\s*[A-Z0-9][A-Z0-9-]{2,}\b", re.I),
    re.compile(r"\b(?:date of birth|DOB)\s*[:#=-]?\s*(?:\d{1,2}[/-]){2}\d{2,4}\b", re.I),
    re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
    re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I),
    re.compile(r"(?<!\w)(?:\+?1[-.\s]?)?(?:\(\d{3}\)|\d{3})[-.\s]?\d{3}[-.\s]?\d{4}(?!\w)"),
)
_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9'-]*")


def redact_identifiers(text: str) -> tuple[str, int]:
    """Mask common direct identifiers before indexing or returning evidence."""
    replacements = 0
    cleaned = text
    for pattern in _PII_PATTERNS:
        cleaned, count = pattern.subn("[REDACTED IDENTIFIER]", cleaned)
        replacements += count
    return cleaned, replacements


def contains_identifier(text: str) -> bool:
    return any(pattern.search(text) for pattern in _PII_PATTERNS)


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_file(path: Path) -> list[tuple[str, str, dict[str, Any]]]:
    """Return (text, locator, extra metadata) records, preserving table rows/pages."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return [(_read_text_file(path), "document", {})]

    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            return [
                (
                    " | ".join(f"{key}: {value}" for key, value in row.items() if key and value),
                    f"row {row_number}",
                    {"table": path.stem, "row_number": row_number, "columns": headers},
                )
                for row_number, row in enumerate(reader, start=2)
            ]

    if suffix == ".json":
        payload = json.loads(_read_text_file(path))
        rows = payload if isinstance(payload, list) else payload.get("records", [payload])
        if not isinstance(rows, list):
            rows = [rows]
        return [
            (json.dumps(row, ensure_ascii=True, sort_keys=True), f"record {index}", {"record_number": index})
            for index, row in enumerate(rows, start=1)
        ]

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel ingestion requires openpyxl. Install requirements.txt.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        records = []
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or f"column_{i + 1}") for i, value in enumerate(next(rows, ())) ]
            for row_number, values in enumerate(rows, start=2):
                cells = {headers[i]: str(value) for i, value in enumerate(values) if i < len(headers) and value is not None}
                if cells:
                    records.append((" | ".join(f"{key}: {value}" for key, value in cells.items()), f"{sheet.title}, row {row_number}", {"table": sheet.title, "row_number": row_number, "columns": headers}))
        workbook.close()
        return records

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF ingestion requires pypdf. Install requirements.txt.") from exc
        return [(page.extract_text() or "", f"page {index}", {"page_number": index}) for index, page in enumerate(PdfReader(path).pages, start=1)]

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Word ingestion requires python-docx. Install requirements.txt.") from exc
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        return [("\n".join(paragraphs + tables), "document", {"table_rows": len(tables)})]

    raise ValueError(f"Unsupported corpus file type: {suffix}")


def _split_text(text: str, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    cleaned = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(cleaned) <= max_chars:
        return [cleaned] if cleaned else []

    parts: list[str] = []
    cursor = 0
    while cursor < len(cleaned):
        end = min(cursor + max_chars, len(cleaned))
        if end < len(cleaned):
            boundary = max(cleaned.rfind("\n", cursor, end), cleaned.rfind(". ", cursor, end), cleaned.rfind(" ", cursor, end))
            if boundary > cursor + max_chars // 2:
                end = boundary + (1 if cleaned[boundary] == "." else 0)
        piece = cleaned[cursor:end].strip()
        if piece:
            parts.append(piece)
        if end >= len(cleaned):
            break
        cursor = max(end - CHUNK_OVERLAP_CHARS, cursor + 1)
    return parts


def _markdown_sections(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, list[str]]] = []
    current = "Overview"
    lines: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^#{1,4}\s+(.+?)\s*$", line)
        if match:
            if lines:
                sections.append((current, lines))
            current = match.group(1)
            lines = []
        else:
            lines.append(line)
    if lines:
        sections.append((current, lines))
    return [(heading, "\n".join(content).strip()) for heading, content in sections if "\n".join(content).strip()]


def load_corpus(corpus_dir: Path | str | None = None) -> dict[str, Any]:
    """Parse every manifest entry into stable, passage-level chunks."""
    root = Path(corpus_dir) if corpus_dir else DEFAULT_CORPUS_DIR
    manifest = json.loads((root / "corpus_manifest.json").read_text(encoding="utf-8"))
    documents: list[dict[str, Any]] = []
    chunks: list[dict[str, Any]] = []
    redactions = 0

    for entry in manifest.get("documents", []):
        source_path = (root / entry["path"]).resolve()
        if root.resolve() not in source_path.parents:
            raise ValueError(f"Corpus path escapes its root: {entry['path']}")
        if not source_path.is_file():
            raise FileNotFoundError(f"Corpus source is missing: {entry['path']}")

        doc = {**entry, "sha256": hashlib.sha256(source_path.read_bytes()).hexdigest()}
        documents.append(doc)
        pieces = _extract_file(source_path)
        chunk_index = 0
        for raw_text, locator, extra in pieces:
            if source_path.suffix.lower() == ".md":
                section_parts = _markdown_sections(raw_text)
            else:
                section_parts = [(locator, raw_text)]
            for section, body in section_parts:
                safe_text, count = redact_identifiers(body)
                redactions += count
                for fragment in _split_text(safe_text):
                    if not fragment:
                        continue
                    chunk_index += 1
                    stable = f"{entry['doc_id']}|{section}|{chunk_index}|{fragment}"
                    chunk_id = "C-" + hashlib.sha1(stable.encode("utf-8")).hexdigest()[:10].upper()
                    chunks.append({
                        "id": chunk_id,
                        "doc_id": entry["doc_id"],
                        "content": fragment,
                        "metadata": {
                            **{key: value for key, value in entry.items() if key not in {"path"}},
                            **extra,
                            "section": section,
                            "source_path": entry["path"],
                            "chunk_number": chunk_index,
                        },
                    })

    return {"documents": documents, "chunks": chunks, "redactions": redactions, "corpus_name": manifest.get("name", "Knowledge corpus")}


if __name__ == "__main__":
    result = load_corpus()
    print(f"Loaded {len(result['documents'])} documents into {len(result['chunks'])} passages; redacted {result['redactions']} identifiers.")
